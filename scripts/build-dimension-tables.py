"""Rebuild bukanpipe-dimension-tables.json from the factory's own spec sheet.

The three pressure tables were first transcribed from a scan of the catalogue
PDF. The factory later supplied research/Bukan_Pipe_dimension-tables.xlsx, a
machine-readable sheet headed "مشخصات لوله های ... تولیدی شرکت لوله پلی اتیلن
بوکان", so that file governs: it is the same company, newer, and needs no OCR.

It confirmed 273 of the 275 water-supply cells. Where it disagreed it was the
scan that was wrong — most importantly the gas table, where the site had been
publishing SDR 11 wall thicknesses for DN 110-225 that this sheet marks as not
produced.

The packaging and coil tables come from catalogue pages the sheet does not
cover, so they are carried through untouched.

Usage: python3 scripts/build-dimension-tables.py
"""

import json
import pathlib
import re

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent.parent
SHEET = ROOT / "research" / "Bukan_Pipe_dimension-tables.xlsx"
TARGET = ROOT / "bukanpipe-dimension-tables.json"


def number(value):
    """A cell as a float, or None for the dash that means 'not produced'."""
    if value in (None, "", "-", "—"):
        return None
    if isinstance(value, str):
        value = value.replace("،", ".").replace(",", ".").strip()
        if value in ("-", "—", ""):
            return None
    return float(value)


def tidy(value):
    return int(value) if value is not None and value == int(value) else value


def wall_with_tolerance(cell):
    """The sheet writes a wall of 3 mm with a +0.4 tolerance as "(0.4+)3"."""
    if cell in (None, "", "-", "—"):
        return None
    match = re.match(r"\(([\d.]+)\+\)([\d.]+)", str(cell).strip())
    if not match:
        return {"min": tidy(number(cell)), "tolerancePlus": 0}
    return {"min": tidy(float(match.group(2))), "tolerancePlus": float(match.group(1))}


def packing(cell):
    """"شاخه 12 متری/کلاف 50 متری" -> both package types and both lengths."""
    text = str(cell or "")
    types, lengths = [], []
    for kind, label in (("coil", "کلاف"), ("branch", "شاخه")):
        for length in re.findall(rf"{label}\s*(\d+)\s*متری", text):
            types.append(kind)
            lengths.append(length)
    return types, lengths


def strip_width(cell):
    if cell in (None, ""):
        return None
    nums = re.findall(r"\d+", str(cell))
    return " - ".join(nums) if nums else None


book = openpyxl.load_workbook(SHEET, data_only=True)
doc = json.loads(TARGET.read_text(encoding="utf-8"))
tables = doc["tables"]

# ---------------------------------------------------------------- water (HD)
hd = book["HD"]
sdr_row, pn80_row, pn100_row, first_data_row = 7, 5, 6, 9
columns, water_rows = [], []
for col in range(3, 14):
    columns.append({
        "sdr": tidy(number(hd.cell(sdr_row, col).value)),
        "pnPe80": tidy(number(str(hd.cell(pn80_row, col).value).replace("PN", ""))),
        "pnPe100": tidy(number(str(hd.cell(pn100_row, col).value).replace("PN", ""))),
    })
for row in range(first_data_row, hd.max_row + 1):
    dn = number(hd.cell(row, 2).value)
    if dn is None:
        continue
    water_rows.append({
        "dnMm": tidy(dn),
        "wallThicknessMm": [tidy(number(hd.cell(row, c).value)) for c in range(3, 14)],
    })

water = tables["waterSupply"]
water["standard"] = "DIN 8074 / ISO 4427-2"
water["nationalStandard"] = "INSO 14427-2"
water["columns"] = columns
water["rows"] = water_rows

# ------------------------------------------------------------------ gas (IGS)
gas_sheet = book["گاز igs"]
gas_rows = []
for row in range(5, gas_sheet.max_row + 1):
    dn = number(gas_sheet.cell(row, 2).value)
    if dn is None:
        continue
    types, lengths = packing(gas_sheet.cell(row, 7).value)
    gas_rows.append({
        "nominalSizeMm": tidy(dn),
        "maxExternalDiameterMm": tidy(number(gas_sheet.cell(row, 3).value)),
        "maxDoubleWidthBranchMm": tidy(number(gas_sheet.cell(row, 4).value)),
        "eMinSdr11Mm": wall_with_tolerance(gas_sheet.cell(row, 6).value),
        "eMinSdr13_6Mm": wall_with_tolerance(gas_sheet.cell(row, 5).value),
        "packageTypes": types,
        "packingLengthM": lengths,
        "yellowBarWidthDepthMm": strip_width(gas_sheet.cell(row, 8).value),
    })

gas = tables["gasSupply"]
gas["standard"] = "ISO 4437 / IGS-M-PL-014-1(04)"
gas["yellowBarDepthNote"] = "حداکثر ۱۰٪ ضخامت جداره"
gas["rows"] = gas_rows

# ----------------------------------------------------------- irrigation (LD)
ld = book["LD"]
irrigation_rows = []
for row in range(8, ld.max_row + 1):
    dn = number(ld.cell(row, 2).value)
    if dn is None:
        continue
    mins, maxes = [], []
    for col in range(3, 15, 2):
        mins.append(tidy(number(ld.cell(row, col).value)))
        maxes.append(tidy(number(ld.cell(row, col + 1).value)))
    irrigation_rows.append({
        "nominalSizeMm": tidy(dn),
        "wallThicknessMm": mins,
        "wallThicknessMaxMm": maxes,
    })

irrigation = tables["dripIrrigation"]
irrigation["standard"] = "ISO 8779"
irrigation["nationalStandard"] = "INSO 7607"
irrigation["rows"] = irrigation_rows

doc["source"] = {
    "document": "Bukan Pipe factory specification sheet (XLSX), with the company catalogue (PDF) behind the packaging tables",
    "sheet": "research/Bukan_Pipe_dimension-tables.xlsx",
    "pages": {"packaging": 13, "coilDimensions": 17},
    "transcribedOn": "2026-09-24",
    "note": (
        "Pressure tables read straight from the factory's own spreadsheet; nothing is "
        "computed, rounded or inferred, and a dash in the sheet is null here. The "
        "packaging and coil tables still come from the catalogue PDF, which the "
        "spreadsheet does not cover."
    ),
}

TARGET.write_text(json.dumps(doc, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(f"water {len(water_rows)} rows · gas {len(gas_rows)} rows · irrigation {len(irrigation_rows)} rows")
